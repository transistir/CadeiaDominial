"""
Issue #179 — export em Excel consolidado por Terra Indígena: um único
arquivo com a cadeia dominial completa de TODOS os imóveis de uma TI, uma
seção por imóvel, no mesmo layout do export por imóvel (issue #50) e do PDF
completo (`cadeia_completa_pdf.html`).

Reaproveita o padrão de `test_issue_166_cri_export.py`: `RequestFactory` +
`SimpleNamespace(is_authenticated=True)` no `request.user` e
`view.__wrapped__(...)` para chamar a view diretamente, pulando o decorator
`login_required` (que exigiria uma sessão de login real).

Decisões de fixture:
- As matrículas dos 3 imóveis do cenário principal são criadas fora de
  ordem (M300, M100, M200) de propósito: se a view devolvesse os imóveis na
  ordem de criação/pk em vez de ordenar por matrícula, o teste de ordenação
  abaixo pegaria isso.
- Uma segunda TI, com um imóvel de matrícula/nome sentinela
  (`M999OUTRATI` / `IMOVEL_DE_OUTRA_TI_179`), prova que a exportação de uma
  TI não vaza dados de outra.
- `Lancamento.objects.bulk_create` evita o signal
  `processar_origens_automaticas_signal` (que criaria documentos de origem
  automaticamente a partir do texto de `origem`), mesmo padrão já usado em
  `test_exportacao_cadeia.py` e `test_issue_166_cri_export.py`.
- `cache.clear()` no `setUp`: o cache de tronco principal por imovel_id é
  compartilhado entre métodos de teste no mesmo processo.
"""

from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.conf import settings
from django.core.cache import cache
from django.db import connection
from django.test import RequestFactory, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from dominial.models import (
    Cartorios,
    Documento,
    DocumentoTipo,
    Imovel,
    Lancamento,
    LancamentoPessoa,
    LancamentoTipo,
    Pessoas,
    TIs,
)
from dominial.services.exportacao_excel_service import CABECALHOS_DETALHADOS
from dominial.views import cadeia_dominial_views


class ExportacaoTisXlsConsolidadoTest(TestCase):
    """
    Cenário principal: uma TI com 3 imóveis (cada um com um documento e um
    lançamento) e uma segunda TI usada apenas para provar isolamento — nada
    dela deve aparecer na planilha da primeira.
    """

    def setUp(self):
        cache.clear()
        self.factory = RequestFactory()

        self.tis = TIs.objects.create(
            nome="TI Teste 179", codigo="TI179", etnia="Teste"
        )
        self.cartorio = Cartorios.objects.create(
            nome="Cartório de Registro de Imóveis Teste 179",
            cns="179001",
            cidade="Cidade 179",
            estado="TS",
        )
        self.proprietario = Pessoas.objects.create(
            nome="Proprietário 179", cpf="10120120312"
        )

        self.tipo_matricula = DocumentoTipo.objects.create(tipo="matricula")
        self.tipo_inicio = LancamentoTipo.objects.create(tipo="inicio_matricula")

        # Matrículas propositalmente fora da ordem de criação.
        self.imovel_m300 = self._criar_imovel_com_cadeia(self.tis, "M300")
        self.imovel_m100 = self._criar_imovel_com_cadeia(self.tis, "M100")
        self.imovel_m200 = self._criar_imovel_com_cadeia(self.tis, "M200")

        # Segunda TI: nada dela deve vazar para a planilha da primeira.
        self.tis_outra = TIs.objects.create(
            nome="TI Outra 179", codigo="TI179OUTRA", etnia="Teste"
        )
        self.imovel_outra_ti = self._criar_imovel_com_cadeia(
            self.tis_outra, "M999OUTRATI", nome="IMOVEL_DE_OUTRA_TI_179"
        )

    def _criar_imovel_com_cadeia(self, tis, matricula, nome=None):
        imovel = Imovel.objects.create(
            terra_indigena_id=tis,
            nome=nome or f"Imóvel {matricula}",
            proprietario=self.proprietario,
            matricula=matricula,
            cartorio=self.cartorio,
        )
        documento = Documento.objects.create(
            imovel=imovel,
            tipo=self.tipo_matricula,
            numero=matricula,
            data=timezone.now().date(),
            cartorio=self.cartorio,
            livro="1",
            folha="1",
        )
        # bulk_create evita o signal de processamento automático de origens
        # (mesmo padrão de test_issue_166_cri_export.py).
        Lancamento.objects.bulk_create([
            Lancamento(
                documento=documento,
                tipo=self.tipo_inicio,
                data=timezone.now().date(),
                cartorio_origem=self.cartorio,
                origem="",
                area=Decimal("0"),
            ),
        ])
        return imovel

    def _request(self, path):
        request = self.factory.get(path)
        request.user = SimpleNamespace(is_authenticated=True)
        return request

    def _exportar(self, tis):
        response = cadeia_dominial_views.exportar_cadeia_dominial_excel_tis.__wrapped__(
            self._request("/excel-tis/"), tis.id
        )
        self.assertEqual(response.status_code, 200)
        return response

    @staticmethod
    def _matriculas_das_secoes(ws):
        """
        Extrai, na ordem em que aparecem na coluna A, as matrículas das
        linhas de seção "IMÓVEL: <matrícula> — <nome>".
        """
        matriculas = []
        for cell in ws["A"]:
            valor = cell.value
            if isinstance(valor, str) and valor.startswith("IMÓVEL: "):
                resto = valor[len("IMÓVEL: "):]
                matriculas.append(resto.split(" — ")[0].strip())
        return matriculas

    # 1 e 2 -------------------------------------------------------------

    def test_rota_responde_200_com_xlsx_e_abre_no_openpyxl(self):
        response = self._exportar(self.tis)

        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", response["Content-Disposition"])

        # Não deve levantar exceção ao abrir.
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("Cadeia Dominial Consolidada", workbook.sheetnames)

    # 3 -------------------------------------------------------------------

    def test_contem_as_tres_secoes_dos_imoveis_da_ti(self):
        response = self._exportar(self.tis)
        ws = load_workbook(BytesIO(response.content)).active

        matriculas = self._matriculas_das_secoes(ws)
        self.assertEqual(len(matriculas), 3)
        self.assertEqual(set(matriculas), {"M100", "M200", "M300"})

    # 4 -------------------------------------------------------------------

    def test_nao_vaza_imovel_de_outra_ti(self):
        response = self._exportar(self.tis)
        ws = load_workbook(BytesIO(response.content)).active

        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str):
                    self.assertNotIn("M999OUTRATI", celula.value)
                    self.assertNotIn("IMOVEL_DE_OUTRA_TI_179", celula.value)

    # 5 -------------------------------------------------------------------

    def test_ordem_das_secoes_segue_a_matricula(self):
        response = self._exportar(self.tis)
        ws = load_workbook(BytesIO(response.content)).active

        self.assertEqual(
            self._matriculas_das_secoes(ws), ["M100", "M200", "M300"]
        )

    def test_erro_parcial_no_imovel_intermediario_nao_sobrepoe_o_seguinte(self):
        renderer_real = cadeia_dominial_views.escrever_secao_documentos
        chamadas = 0

        def renderer_com_falha_parcial(
            ws, cadeia_completa, linha_inicial, estilos
        ):
            nonlocal chamadas
            chamadas += 1
            if chamadas == 2:
                # Simula uma exceção depois de o renderer já ter avançado e
                # escrito parte da seção do segundo imóvel.
                ws.cell(
                    row=linha_inicial + 3,
                    column=1,
                    value="LINHA PARCIAL DO IMÓVEL M200",
                )
                raise ValueError("falha de fixture no renderer")
            return renderer_real(
                ws, cadeia_completa, linha_inicial, estilos
            )

        with (
            patch.object(
                cadeia_dominial_views,
                "escrever_secao_documentos",
                side_effect=renderer_com_falha_parcial,
            ),
            patch.object(cadeia_dominial_views.logger, "exception"),
        ):
            response = self._exportar(self.tis)

        ws = load_workbook(BytesIO(response.content)).active
        linhas_por_valor = {
            cell.value: cell.row
            for cell in ws["A"]
            if isinstance(cell.value, str)
        }

        self.assertLess(
            linhas_por_valor["LINHA PARCIAL DO IMÓVEL M200"],
            linhas_por_valor["Erro ao exportar este imóvel."],
        )
        self.assertLess(
            linhas_por_valor["Erro ao exportar este imóvel."],
            linhas_por_valor["IMÓVEL: M300 — Imóvel M300"],
        )
        self.assertEqual(
            self._matriculas_das_secoes(ws), ["M100", "M200", "M300"]
        )
        self.assertIn("Matrícula: M300", linhas_por_valor)

    def test_queries_ficam_sob_teto_com_tres_imoveis_e_dois_documentos(self):
        tipo_transcricao = DocumentoTipo.objects.create(tipo="transcricao")
        pessoa_transmitente = Pessoas.objects.create(
            nome="Transmitente 179", cpf="11122233344"
        )
        pessoa_adquirente = Pessoas.objects.create(
            nome="Adquirente 179", cpf="55566677788"
        )

        for indice, imovel in enumerate(
            (self.imovel_m100, self.imovel_m200, self.imovel_m300), start=1
        ):
            numero_origem = f"T{indice}79"
            documento_origem = Documento.objects.create(
                imovel=imovel,
                tipo=tipo_transcricao,
                numero=numero_origem,
                data=timezone.now().date(),
                cartorio=self.cartorio,
                livro="2",
                folha="2",
            )
            lancamento_origem = Lancamento.objects.create(
                documento=documento_origem,
                tipo=self.tipo_inicio,
                data=timezone.now().date(),
                cartorio_origem=self.cartorio,
                origem="",
            )
            lancamento_principal = Lancamento.objects.get(
                documento__imovel=imovel,
                documento__tipo=self.tipo_matricula,
            )
            lancamento_principal.origem = numero_origem
            lancamento_principal.save(update_fields=["origem"])

            LancamentoPessoa.objects.bulk_create(
                [
                    LancamentoPessoa(
                        lancamento=lancamento_principal,
                        pessoa=pessoa_transmitente,
                        tipo="transmitente",
                    ),
                    LancamentoPessoa(
                        lancamento=lancamento_principal,
                        pessoa=pessoa_adquirente,
                        tipo="adquirente",
                    ),
                    LancamentoPessoa(
                        lancamento=lancamento_origem,
                        pessoa=pessoa_transmitente,
                        tipo="transmitente",
                    ),
                    LancamentoPessoa(
                        lancamento=lancamento_origem,
                        pessoa=pessoa_adquirente,
                        tipo="adquirente",
                    ),
                ]
            )

        cache.clear()
        with CaptureQueriesContext(connection) as queries:
            response = self._exportar(self.tis)

        # A recursão histórica do serviço de cadeia ainda custa consultas por
        # nó (fora do escopo da #179). O teto cobre 3 imóveis x 2 documentos
        # com folga pequena, mas não admite as 2 consultas adicionais por
        # lançamento que `.transmitentes/.adquirentes.filter()` causavam.
        self.assertLessEqual(len(queries), 100)
        ws = load_workbook(BytesIO(response.content)).active
        valores = [cell.value for row in ws.iter_rows() for cell in row]
        self.assertIn("Transmitente 179", valores)
        self.assertIn("Adquirente 179", valores)

    # 6 -------------------------------------------------------------------

    def test_estrutura_de_colunas_igual_ao_export_por_imovel(self):
        response = self._exportar(self.tis)
        ws = load_workbook(BytesIO(response.content)).active

        header = None
        for linha in ws.iter_rows():
            if linha and linha[0].value == "Nº":
                header = linha
                break
        self.assertIsNotNone(
            header, "Linha de cabeçalho detalhado (começando em 'Nº') não encontrada"
        )

        valores = [cell.value for cell in header]
        self.assertEqual(valores, CABECALHOS_DETALHADOS)
        self.assertEqual(valores[3], "CRI")
        self.assertEqual(valores[9], "CRI")

    # 7 -------------------------------------------------------------------

    def test_nao_menciona_tronco_principal_ou_secundario(self):
        response = self._exportar(self.tis)
        ws = load_workbook(BytesIO(response.content)).active

        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str):
                    self.assertNotIn("TRONCO PRINCIPAL", celula.value)
                    self.assertNotIn("TRONCO SECUNDÁRIO", celula.value)

    # 8 -------------------------------------------------------------------

    def test_area_zerada_exibe_traco_na_coluna_14(self):
        # Os lançamentos das fixtures têm area=Decimal("0"); por decisão
        # de produto (Hiure, 10/09/2026) área zerada é exibida como "-",
        # nunca como "0,0000".
        response = self._exportar(self.tis)
        ws = load_workbook(BytesIO(response.content)).active

        valores_coluna_14 = [
            ws.cell(row=linha[0].row, column=14).value for linha in ws.iter_rows()
        ]
        self.assertIn("-", valores_coluna_14)
        self.assertNotIn("0,0000", valores_coluna_14)

    # 10 ------------------------------------------------------------------

    def test_rota_resolve_por_nome(self):
        url = reverse("exportar_cadeia_tis_excel", kwargs={"tis_id": self.tis.id})
        self.assertEqual(url, f"/dominial/tis/{self.tis.id}/imoveis/excel/")

    # 11 ------------------------------------------------------------------

    def test_botao_presente_em_imoveis_html(self):
        template = (
            Path(settings.BASE_DIR) / "templates" / "dominial" / "imoveis.html"
        ).read_text(encoding="utf-8")

        self.assertIn("{% url 'exportar_cadeia_tis_excel' tis_id=tis.id %}", template)
        self.assertIn("Exportar XLS da TI", template)

    def test_botao_presente_em_tis_detail_html(self):
        template = (
            Path(settings.BASE_DIR) / "templates" / "dominial" / "tis_detail.html"
        ).read_text(encoding="utf-8")

        self.assertIn("{% url 'exportar_cadeia_tis_excel' tis_id=tis.id %}", template)
        self.assertIn("Exportar XLS da TI", template)


class ExportacaoTisXlsImovelSemDocumentosTest(TestCase):
    """
    Caso dedicado (item 9 da issue #179): um imóvel sem nenhum documento
    cadastrado ainda deve gerar sua seção (com o aviso "Sem documentos
    cadastrados."), numa TI própria — para não alterar a contagem de 3
    seções exercitada em `ExportacaoTisXlsConsolidadoTest`.
    """

    def setUp(self):
        cache.clear()
        self.factory = RequestFactory()

        self.tis = TIs.objects.create(
            nome="TI Sem Documentos 179", codigo="TI179SD", etnia="Teste"
        )
        self.cartorio = Cartorios.objects.create(
            nome="Cartório Teste 179 SD", cns="179002", cidade="Cidade", estado="TS"
        )
        self.proprietario = Pessoas.objects.create(
            nome="Proprietário 179 SD", cpf="98765432100"
        )
        self.imovel = Imovel.objects.create(
            terra_indigena_id=self.tis,
            nome="Imóvel Sem Documentos",
            proprietario=self.proprietario,
            matricula="M400",
            cartorio=self.cartorio,
        )

    def _request(self, path):
        request = self.factory.get(path)
        request.user = SimpleNamespace(is_authenticated=True)
        return request

    def test_imovel_sem_documentos_gera_secao_com_aviso(self):
        response = cadeia_dominial_views.exportar_cadeia_dominial_excel_tis.__wrapped__(
            self._request("/excel-tis/"), self.tis.id
        )
        self.assertEqual(response.status_code, 200)

        ws = load_workbook(BytesIO(response.content)).active
        valores_coluna_a = [cell.value for cell in ws["A"]]

        self.assertIn("IMÓVEL: M400 — Imóvel Sem Documentos", valores_coluna_a)
        self.assertIn("Sem documentos cadastrados.", valores_coluna_a)
