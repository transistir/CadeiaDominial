"""
Issue #204 — XLS da cadeia dominial sem o rótulo "Importado" e com impressão
em A4 paisagem, como o PDF.

1. Juridicamente não existe "documento importado": documentos vindos de
   outra cadeia são compartilhados entre cadeias. `is_importado`
   (`CadeiaCompletaService`) continua existindo como sinal interno — as
   estatísticas ainda o contam —, mas nenhuma célula do XLS pode conter
   "importado", nem como prefixo do título do documento nem em outro lugar.
2. O XLS imprime como o ``@page { size: A4 landscape; margin: 1.5cm; }`` de
   `cadeia_dominial_pdf.css`: largura em uma página, altura livre
   (``fitToHeight = 0``) e margens de ~1,5 cm. Vale para os dois exports que
   usam `renderizar_planilha_imovel` — `exportar_cadeia_dominial_excel` e
   `exportar_cadeia_dominial_excel_tis` — e, no consolidado, para TODAS as
   abas: "Resumo", abas dos imóveis e aba de imóvel que falhou.

Fixture (mesmo desenho da #201): a matrícula M204 do imóvel A tem como
origem a transcrição T204, que pertence ao imóvel B e também é origem da
matrícula M205 do próprio B. T204 é, portanto, compartilhada entre as duas
cadeias e aparece na de A com `is_importado=True`.
`test_fixture_tem_documento_compartilhado_marcado_como_importado` prova
isso: sem esse caminho, a varredura por "importado" passaria mesmo com o bug.

As views são chamadas via `view.__wrapped__(...)` com `RequestFactory`, como
em `test_issue_179_xls_consolidado_tis.py`; `Lancamento.objects.bulk_create`
evita o signal `processar_origens_automaticas_signal`; `cache.clear()` no
`setUp` porque o tronco principal é cacheado por id de imóvel.
"""

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.cache import cache
from django.test import RequestFactory, TestCase
from openpyxl import load_workbook

from dominial.models import (
    Cartorios,
    Documento,
    DocumentoTipo,
    Imovel,
    Lancamento,
    LancamentoTipo,
    Pessoas,
    TIs,
)
from dominial.services.cadeia_completa_service import CadeiaCompletaService
from dominial.views import cadeia_dominial_views


CONTENT_TYPE_XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
# 1,5 cm em polegadas, a unidade de `PageMargins` no openpyxl.
MARGEM_PDF_POLEGADAS = 1.5 / 2.54


class ExportacaoXlsIssue204Test(TestCase):
    def setUp(self):
        cache.clear()
        self.factory = RequestFactory()

        self.tis = TIs.objects.create(
            nome="TI Teste 204", codigo="TI204", etnia="Teste"
        )
        self.cartorio = Cartorios.objects.create(
            nome="Cartório de Registro de Imóveis Teste 204",
            cns="204001",
            cidade="Cidade 204",
            estado="TS",
        )
        proprietario = Pessoas.objects.create(nome="Proprietário 204")

        tipo_matricula = DocumentoTipo.objects.create(tipo="matricula")
        tipo_transcricao = DocumentoTipo.objects.create(tipo="transcricao")
        tipo_inicio = LancamentoTipo.objects.create(tipo="inicio_matricula")

        self.imovel_a = Imovel.objects.create(
            terra_indigena_id=self.tis,
            nome="Imóvel A 204",
            proprietario=proprietario,
            matricula="M204",
            cartorio=self.cartorio,
        )
        self.imovel_b = Imovel.objects.create(
            terra_indigena_id=self.tis,
            nome="Imóvel B 204",
            proprietario=proprietario,
            matricula="M205",
            cartorio=self.cartorio,
        )

        doc_m204 = Documento.objects.create(
            imovel=self.imovel_a, tipo=tipo_matricula, numero="M204",
            data=date(2020, 1, 10), cartorio=self.cartorio,
            livro="1", folha="1",
        )
        doc_m205 = Documento.objects.create(
            imovel=self.imovel_b, tipo=tipo_matricula, numero="M205",
            data=date(2019, 5, 20), cartorio=self.cartorio,
            livro="1", folha="2",
        )
        self.doc_t204 = Documento.objects.create(
            imovel=self.imovel_b, tipo=tipo_transcricao, numero="T204",
            data=date(2010, 3, 1), cartorio=self.cartorio,
            livro="2", folha="3",
        )
        Lancamento.objects.bulk_create([
            Lancamento(
                documento=doc_m204, tipo=tipo_inicio,
                data=date(2020, 1, 10), cartorio_origem=self.cartorio,
                origem="T204",
            ),
            Lancamento(
                documento=doc_m205, tipo=tipo_inicio,
                data=date(2019, 5, 20), cartorio_origem=self.cartorio,
                origem="T204",
            ),
            Lancamento(
                documento=self.doc_t204, tipo=tipo_inicio,
                data=date(2010, 3, 1), cartorio_origem=self.cartorio,
                origem="",
            ),
        ])

    # -- Helpers ----------------------------------------------------------

    def _request(self, path):
        request = self.factory.get(path)
        request.user = SimpleNamespace(is_authenticated=True)
        return request

    def _abrir_xlsx(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], CONTENT_TYPE_XLSX)
        return load_workbook(BytesIO(response.content))

    def _exportar_individual(self, imovel):
        return self._abrir_xlsx(
            cadeia_dominial_views.exportar_cadeia_dominial_excel.__wrapped__(
                self._request("/excel/"), self.tis.id, imovel.id
            )
        )

    def _exportar_consolidado(self):
        return self._abrir_xlsx(
            cadeia_dominial_views.exportar_cadeia_dominial_excel_tis.__wrapped__(
                self._request("/excel-tis/"), self.tis.id
            )
        )

    def _assert_sem_rotulo_importado(self, workbook):
        ocorrencias = [
            f"{ws.title}!{cell.coordinate}: {cell.value!r}"
            for ws in workbook.worksheets
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
            and "importado" in str(cell.value).casefold()
        ]
        self.assertEqual(ocorrencias, [])

    def _assert_impressao_a4_paisagem(self, ws):
        self.assertEqual(ws.page_setup.orientation, "landscape")
        # `Worksheet.PAPERSIZE_A4` é a string '9', mas `paperSize` é um
        # descritor inteiro no openpyxl: sem o `int()` a comparação seria
        # sempre falsa.
        self.assertEqual(ws.page_setup.paperSize, int(ws.PAPERSIZE_A4))
        self.assertIs(ws.sheet_properties.pageSetUpPr.fitToPage, True)
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        # Lido do arquivo salvo: `None` aqui significaria atributo omitido,
        # que o Excel trata como 1 página de altura.
        self.assertEqual(ws.page_setup.fitToHeight, 0)
        for lado in ("left", "right", "top", "bottom"):
            self.assertAlmostEqual(
                getattr(ws.page_margins, lado), MARGEM_PDF_POLEGADAS, places=3
            )

    # -- Fixture ----------------------------------------------------------

    def test_fixture_tem_documento_compartilhado_marcado_como_importado(self):
        contexto = CadeiaCompletaService().get_cadeia_completa(
            self.tis.id, self.imovel_a.id
        )
        itens_t204 = [
            item
            for tronco in contexto["cadeia_completa"]
            for item in tronco["documentos"]
            if item["documento"].id == self.doc_t204.id
        ]

        self.assertTrue(itens_t204)
        # A #204 remove só a exibição: o sinal interno segue intacto.
        for item in itens_t204:
            self.assertTrue(item["is_importado"])
        self.assertGreaterEqual(
            contexto["estatisticas"]["documentos_importados"], 1
        )

    # -- Export por imóvel ------------------------------------------------

    def test_xls_individual_nao_contem_importado_em_nenhuma_celula(self):
        workbook = self._exportar_individual(self.imovel_a)

        valores_coluna_a = [cell.value for cell in workbook.active["A"]]
        # Antes da #204, T204 saía como "[Importado] Transcrição: T204".
        self.assertIn("Matrícula: M204", valores_coluna_a)
        self.assertIn("Transcrição: T204", valores_coluna_a)
        self._assert_sem_rotulo_importado(workbook)

    def test_xls_individual_imprime_em_a4_paisagem(self):
        workbook = self._exportar_individual(self.imovel_a)

        self.assertEqual(workbook.sheetnames, ["Cadeia Dominial Geral"])
        self._assert_impressao_a4_paisagem(workbook.active)

    # -- Export consolidado por TI ----------------------------------------

    def test_xls_consolidado_nao_contem_importado_em_nenhuma_aba(self):
        workbook = self._exportar_consolidado()

        self.assertEqual(workbook.sheetnames, ["Resumo", "M204", "M205"])
        self.assertIn(
            "Transcrição: T204",
            [cell.value for cell in workbook["M204"]["A"]],
        )
        self._assert_sem_rotulo_importado(workbook)

    def test_xls_consolidado_imprime_em_a4_paisagem_em_todas_as_abas(self):
        workbook = self._exportar_consolidado()

        self.assertEqual(workbook.sheetnames, ["Resumo", "M204", "M205"])
        for ws in workbook.worksheets:
            with self.subTest(aba=ws.title):
                self._assert_impressao_a4_paisagem(ws)

    def test_xls_consolidado_aba_de_imovel_com_erro_imprime_em_a4_paisagem(self):
        # A falha acontece em `get_cadeia_completa`, ANTES de
        # `renderizar_planilha_imovel`: a aba de erro não pode depender do
        # renderer para receber a configuração de impressão.
        get_cadeia_completa_real = CadeiaCompletaService.get_cadeia_completa
        imovel_com_erro_id = self.imovel_a.id

        def falha_no_imovel_a(service, tis_id, imovel_id):
            if imovel_id == imovel_com_erro_id:
                raise ValueError("falha de fixture na cadeia do imóvel A")
            return get_cadeia_completa_real(service, tis_id, imovel_id)

        with (
            patch.object(
                CadeiaCompletaService,
                "get_cadeia_completa",
                autospec=True,
                side_effect=falha_no_imovel_a,
            ),
            patch.object(cadeia_dominial_views.logger, "exception"),
        ):
            workbook = self._exportar_consolidado()

        self.assertIn(
            "Erro ao exportar este imóvel.",
            [cell.value for cell in workbook["M204"]["A"]],
        )
        self.assertIn(
            "Matrícula: M205",
            [cell.value for cell in workbook["M205"]["A"]],
        )
        for ws in workbook.worksheets:
            with self.subTest(aba=ws.title):
                self._assert_impressao_a4_paisagem(ws)
